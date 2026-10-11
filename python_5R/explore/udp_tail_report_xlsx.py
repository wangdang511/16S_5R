import pandas as pd, numpy as np, pickle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SC='/tmp/claude-0/-home-user-16S-5R/0a157df1-6424-5d79-bd19-c9f07612faab/scratchpad/udp/'
X=pd.read_pickle(SC+'X.pkl'); R=pickle.load(open(SC+'udp_eval.pkl','rb')); B=R['base']
base=dict(H=float(np.triu(B['H']).min()),E=float(B['E'].min()),bad=float(B['bad']),hp=float(B['hg'].min()))
f=lambda **k: Font(**{'name':'Arial','size':10,**k})
hf=PatternFill('solid',fgColor='1F3864'); wb=Workbook()
thin=Side(style='thin',color='BFBFBF'); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
fills={'A':PatternFill('solid',fgColor='E2EFDA'),'B':PatternFill('solid',fgColor='FFF2CC'),'C':PatternFill('solid',fgColor='F8CBAD')}
s=wb.active; s.title='说明'
L=[('UDP 内联尾巴对 33 条引物池的二聚体、发夹和脱靶的影响（Illumina UDP Set A，UDP0001–0096）',1),
('做法：每个管里，i5（“Bases in Adapter”那一列的 10 nt）加在所有正向引物 5′端，同编号的 i7 加在所有反向引物 5′端，没有接头；198 个展开序列两两评估，58 °C，primer3（50 mM Na⁺、2 mM Mg²⁺、0.2 mM dNTP、250 nM），再与无尾巴基线比较。我把 UDP0001 的 i5/i7 当作你说的 UDP501/UDP701，以此类推；如果你的编号对应不同，请告诉我。',0),
('基线（无尾巴，58 °C）：全局二聚体最差 ΔG %.1f，3′端锚定最差 %.1f，发夹最差 %.1f kcal/mol，Olivar 总坏度 %.0f。'%(base['H'],base['E'],base['hp'],base['bad']),0),
('分级（阈值是我按基线设的经验阈值，没有文献依据）：A = 全局 ΔG ≥ −7 且 3′端 ≥ −4.5；B = 全局 ≥ −8.5 且 3′端 ≥ −5.5；C = 其余。96 管中 A 17、B 35、C 44。推荐 48 管 = 全部 A + B 中 Olivar 坏度最低的 31 个，仍含 B 级管，最差全局 −8.5、3′端 −5.4。',0),
('结论：UDP 尾巴使二聚体明显变差，主要是尾巴与别的引物特异部分形成互补（例如 UDP0034 的 i7 含 CGCGGCTG，与 A3-F 的 CAGCCGCG 完全互补，全局 ΔG −14.2）。脱靶基本不受影响：尾巴在 5′端，3′端种子不变，全长错配数只会增加。',0),
('脱靶：把尾巴并入后，人和猪基因组全长 ≤2 错配的位点在所有管里都是 0；≤3 错配最多 2 个，且都是基因特异部分本来就 ≤3 错配的位点。',0),
('没有评估的：真实接头（只评估内联 UDP）；管间（不同样本的）串扰；i5 与 i7 的重新配对（可以从更多 UDP 里挑出互补性更低的组合，需要 UDP0097–0384 的序列）。',0),
('数据：96 管完整引物见 docs/primer_design/udp_tail_full_oligos_96tubes.csv；UDP 序列见 docs/primer_design/udp_tail/udp_set_A_1_96.tsv。',0),
('版权：Illumina 序列文件要求发表或分发时注明 “Oligonucleotide sequences © 2025 Illumina, Inc. All rights reserved.”，并限定衍生序列只用于 Illumina 仪器和产品。',0)]
s.column_dimensions['A'].width=140
for i,(t,b) in enumerate(L,1):
    c=s.cell(i,1,t); c.font=f(bold=bool(b),size=12 if b else 10); c.alignment=Alignment(wrap_text=True,vertical='top')
ws=wb.create_sheet('96管汇总')
cols=[('管','k1'),('UDP','name'),('i5（正向引物尾巴）','i5'),('i7（反向引物尾巴）','i7'),('分级','tier'),('推荐48','rec48'),('排名','rank'),
('全局二聚体最差 ΔG','minH'),('最差配对','worstH'),('该配对基线 ΔG','worstH_base'),('ΔG≤−6 的配对数','n6'),('ΔG≤−7 的配对数','n7'),
('3′端最差 ΔG','minE'),('最差 3′配对（前者 3′端延伸）','worstE'),('该配对基线 ΔG ','worstE_base'),('3′ΔG≤−4 的有向配对数','nE4'),('3′ΔG≤−5','nE5'),
('正向×正向最差','H_FF'),('反向×反向最差','H_RR'),('正向×反向最差','H_FR'),
('发夹最差 ΔG','hp_dg'),('发夹最差引物','hp_oligo'),('发夹 ΔG≤−3 的展开数','n_hp3'),('尾巴×尾巴 ΔG','tt_h'),
('Olivar 总坏度','bad'),('人 全长≤3错配位点','o_h3'),('人 全长≤4','o_h4'),('猪 全长≤3错配位点','o_p3'),('猪 全长≤4','o_p4')]
X['k1']=X.k+1
for j,(h,_) in enumerate(cols,1):
    c=ws.cell(1,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=hf; c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center'); c.border=bd
for i,r in enumerate(X.sort_values('k').itertuples(),2):
    for j,(h,a) in enumerate(cols,1):
        v=getattr(r,a)
        if isinstance(v,(np.floating,float)): v=round(float(v)) if a=='bad' else round(float(v),2)
        if isinstance(v,np.integer): v=int(v)
        c=ws.cell(i,j,v); c.font=f(); c.border=bd
    ws.cell(i,5).fill=fills[r.tier]
for j,w in enumerate([5,12,14,14,6,8,6,10,34,10,9,9,10,40,10,10,8,9,9,9,9,18,10,9,12,9,8,9,8],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.row_dimensions[1].height=58; ws.freeze_panes='C2'
st=wb.create_sheet('统计'); st['A1']='项目'; st['B1']='数量'
for c in ('A1','B1'): st[c].font=f(bold=True)
items=[('A 级管',"=COUNTIF('96管汇总'!E2:E97,\"A\")"),('B 级管',"=COUNTIF('96管汇总'!E2:E97,\"B\")"),('C 级管',"=COUNTIF('96管汇总'!E2:E97,\"C\")"),
('推荐 48 管',"=COUNTIF('96管汇总'!F2:F97,\"推荐\")"),('推荐管中 A 级',"=COUNTIFS('96管汇总'!F2:F97,\"推荐\",'96管汇总'!E2:E97,\"A\")"),
('全部管全局 ΔG ≤ −8 的数量',"=COUNTIF('96管汇总'!H2:H97,\"<=-8\")"),('全部管 3′ΔG ≤ −5 的数量',"=COUNTIF('96管汇总'!M2:M97,\"<=-5\")"),
('基线全局最差 ΔG',round(base['H'],2)),('基线 3′最差 ΔG',round(base['E'],2)),('基线发夹最差 ΔG',round(base['hp'],2)),('基线 Olivar 坏度',round(base['bad']))]
for i,(a,b) in enumerate(items,2):
    st.cell(i,1,a).font=f(); st.cell(i,2,b).font=f()
st.column_dimensions['A'].width=32
w2=wb.create_sheet('推荐48管'); r48=X[X.rec48=='推荐'].sort_values('rank')
for j,h in enumerate(['排名','UDP','i5','i7','分级','全局最差 ΔG','3′最差 ΔG','Olivar 坏度'],1):
    c=w2.cell(1,j,h); c.font=f(bold=True,color='FFFFFF'); c.fill=hf; c.border=bd
for i,r in enumerate(r48.itertuples(),2):
    for j,v in enumerate([r.rank,r.name,r.i5,r.i7,r.tier,round(r.minH,2),round(r.minE,2),round(r.bad)],1):
        c=w2.cell(i,j,v); c.font=f(); c.border=bd
    w2.cell(i,5).fill=fills[r.tier]
for j,w in enumerate([6,12,14,14,6,12,10,12],1): w2.column_dimensions[get_column_letter(j)].width=w
wb.save('/home/user/16S_5R/docs/primer_design/udp_tail_evaluation_pool33_v3.xlsx')
print(','.join(r48.name.tolist()))
