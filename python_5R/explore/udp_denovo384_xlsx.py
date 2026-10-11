import sys, json, csv
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
fin = json.load(open(sys.argv[1])); out = sys.argv[2]; csvout = sys.argv[3]
i5, i7, res = fin['i5'], fin['i7'], fin['res']
f = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k}); hf = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); s = wb.active; s.title = '说明'
L = [('自建 UDP：384 对 i5/i7（10 nt 内联索引，4 个 96 组块，每块 2 个 48 组）', 1),
('生成方式：程序从头生成，不借用 Illumina 序列；与 Illumina UDP0001–0384 的 i5/i7 没有任何相同序列（含反向互补），最小汉明距离 1。', 0),
('单条规则：GC 4–6/10；无同聚物 ≥3、无二核苷酸重复 ≥3 次、无 GGGG；前两位不同时为 G；不含 TruSeq/Nextera/P5/P7 接头的 6-mer；非回文（与自身反向互补的编辑距离 ≥4）。', 0),
('集合规则：768 条序列（384 个 i5 + 384 个 i7）两两 Levenshtein 编辑距离 ≥4（实测最小 4）；与对方反向互补的编辑距离 ≥3（严格要求 ≥4 时只能凑出 464 条，不够 768，因此放宽到 3，其余不变）。', 0),
('每个周期（第 1–10 位）：GC、A+C、A+T 占比 40–60%，且 A/C/G/T 各 ≥15%。检验范围：每个 48 组、每个 96 组块（i5、i7 各自）和全部 384 个，共 26 组，全部达标（见“每周期统计”，公式计算）。', 0),
('配对：同一块里 i5 与 i7 按编号配对（DN001 的 i5 与 i7 在同一管）。配对只是占位，你们结合引物筛选后可以重新配对（i5 之间、i7 之间的约束不受配对影响）。', 0),
('二聚体预筛列 H、E：该索引（作为 5′ 尾巴的 10 nt）与 33 条引物池 198 个展开序列（不带尾巴）在 58 °C 的最差异源二聚体 ΔG（H）和池中引物 3′ 端对它的最差 3′ 锚定 ΔG（E），只作排序参考；每个管（i5+i7 加到整池后）的完整二聚体评估在后台运行，结果会补充。', 0),
('没有评估的：真实接头；实验验证（条形码紧邻引物会带来扩增偏好，建议做同样品多条形码的对照实验）。', 0)]
s.column_dimensions['A'].width = 150
for i, (t, b) in enumerate(L, 1):
    c = s.cell(i, 1, t); c.font = f(bold=bool(b), size=12 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
w = wb.create_sheet('384对'); H5 = fin['H'][:0]
heads = ['配对', '96 块', '48 组', 'i5 序列', 'i7 序列', 'i5 GC', 'i7 GC', 'i5 预筛 H', 'i5 预筛 E', 'i7 预筛 H', 'i7 预筛 E']
for j, h in enumerate(heads, 1):
    c = w.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(horizontal='center', wrap_text=True); c.border = bd
for q in range(10): w.cell(1, 13 + q, f'i5 第{q + 1}位'); w.cell(1, 23 + q, f'i7 第{q + 1}位')
sl = None
HE = fin['H']; EE = fin['E']
# slots 顺序：groups 0-7 = i5 半块，8-15 = i7 半块；fin['H'] 与之对应
for k in range(384):
    r = k + 2; blk = k // 96 + 1; half = k // 48 + 1
    vals = [f'DN{k + 1:03d}', blk, half, i5[k], i7[k], None, None, round(HE[k], 2), round(EE[k], 2), round(HE[384 + k], 2), round(EE[384 + k], 2)]
    for j, v in enumerate(vals, 1):
        c = w.cell(r, j, v); c.font = f(color='0000FF' if j in (4, 5) else '000000'); c.border = bd
    w.cell(r, 6, f'=LEN(SUBSTITUTE(SUBSTITUTE(D{r},"A",""),"T",""))'); w.cell(r, 7, f'=LEN(SUBSTITUTE(SUBSTITUTE(E{r},"A",""),"T",""))')
    for q in range(10): w.cell(r, 13 + q, f'=MID($D{r},{q + 1},1)'); w.cell(r, 23 + q, f'=MID($E{r},{q + 1},1)')
for j, wd in enumerate([8, 6, 6, 14, 14, 6, 6, 9, 9, 9, 9], 1): w.column_dimensions[get_column_letter(j)].width = wd
w.freeze_panes = 'D2'; w.row_dimensions[1].height = 32
c1 = wb.create_sheet('每周期统计')
for j, h in enumerate(['范围', '序列', '周期', 'A', 'C', 'G', 'T', 'GC%', 'A+C%', 'A+T%', '是否达标'], 1):
    c = c1.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf
row = 2
specs = [(f'48 组 {g + 1}', g * 48 + 2, (g + 1) * 48 + 1, 48) for g in range(8)] + [(f'96 块 {b + 1}', b * 96 + 2, (b + 1) * 96 + 1, 96) for b in range(4)] + [('全部 384', 2, 385, 384)]
for lab, r0, r1, N in specs:
    for name, base in (('i5', 13), ('i7', 23)):
        for p in range(10):
            hc = get_column_letter(base + p); rng = f"'384对'!${hc}${r0}:${hc}${r1}"
            c1.cell(row, 1, lab); c1.cell(row, 2, name); c1.cell(row, 3, p + 1)
            for j, b in enumerate('ACGT', 4): c1.cell(row, j, f'=COUNTIF({rng},"{b}")')
            c1.cell(row, 8, f'=(E{row}+F{row})/{N}'); c1.cell(row, 9, f'=(D{row}+E{row})/{N}'); c1.cell(row, 10, f'=(D{row}+G{row})/{N}')
            c1.cell(row, 11, f'=IF(AND(H{row}>=0.4,H{row}<=0.6,I{row}>=0.4,I{row}<=0.6,J{row}>=0.4,J{row}<=0.6,MIN(D{row}:G{row})>={round(0.15 * N, 2)}),"是","否")')
            for j in (8, 9, 10): c1.cell(row, j).number_format = '0.0%'
            row += 1
for r in c1.iter_rows(min_row=2):
    for c in r: c.font = f()
wb.save(out)
with open(csvout, 'w', newline='', encoding='utf-8-sig') as fh:
    cw = csv.writer(fh); cw.writerow(['配对', '96块', '48组', 'i5', 'i7'])
    for k in range(384): cw.writerow([f'DN{k + 1:03d}', k // 96 + 1, k // 48 + 1, i5[k], i7[k]])
print('ok')
