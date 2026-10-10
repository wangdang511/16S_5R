"""导出重新配对后的 96 对（前 48 对 = 核心 48 组）：Excel（含独立验证和每周期公式统计）+ 完整引物 csv。
分级优先用整管评估的实测结果（eval pkl 给出时），否则用热点代理的预测。
用法：python3 udp960_export.py <pairs96 json> <out xlsx> <out csv> [eval pkl ...]"""
import sys, json, pickle, collections, re, csv, itertools
import numpy as np, openpyxl
from rapidfuzz.process import cdist
from rapidfuzz.distance import Levenshtein
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
pj, outx, outc = sys.argv[1:4]; pkls = sys.argv[4:]
D = '/home/user/16S_5R/docs/primer_design/'
pairs = json.load(open(pj))['pairs']; rc = lambda s: s[::-1].translate(str.maketrans('ACGT', 'TGCA'))
meas = {}
for p in pkls:
    r = pickle.load(open(p, 'rb')); N0 = len(r['EX']); iu = np.triu_indices(N0)
    for t in r['tubes'].values():
        h = float(t['H'][iu].min()); e = float(t['E'].min())
        meas[t['name']] = dict(minH=h, minE=e, hp=float(t['hg'].min()), bad=float(t['bad']), tier='A' if (h >= -7 and e >= -4.5) else 'B' if (h >= -8.5 and e >= -5.5) else 'C')
def info(k):
    nm = f'P{k + 1:03d}'
    if nm in meas: m = meas[nm]; return dict(tier=m['tier'], H=m['minH'], E=m['minE'], hp=m['hp'], bad=m['bad'], src='实测')
    p = pairs[k]; return dict(tier=p['tier'], H=p['H'], E=p['E'], hp=None, bad=None, src='预测')
i5 = [p['seq5'] for p in pairs]; i7 = [p['seq7'] for p in pairs]
def verify(idx):
    a = [i5[k] for k in idx]; b = [i7[k] for k in idx]; S = a + b; n = len(S); N = len(idx)
    Dm = cdist(S, S, scorer=Levenshtein.distance, dtype=np.uint8, workers=2); np.fill_diagonal(Dm, 99)
    cross = cdist(a, [rc(s) for s in b], scorer=Levenshtein.distance, dtype=np.uint8, workers=2); cross2 = cdist([rc(s) for s in a], b, scorer=Levenshtein.distance, dtype=np.uint8, workers=2)
    own = min(Levenshtein.distance(s, rc(s)) for s in S); ok = True
    for M in (a, b):
        for p in range(10):
            c = collections.Counter(s[p] for s in M); v = (c['G'] + c['C'], c['A'] + c['C'], c['A'] + c['T'])
            ok &= all(0.4 * N <= x <= 0.6 * N for x in v) and min(c[x] for x in 'ACGT') >= 0.15 * N
    ill = set()
    for l in open(D + 'udp_tail/udp_all_1_384.tsv'):
        q = l.split(); ill |= {q[1], q[3]}
    return {'最小编辑距离（全部序列）': int(Dm.min()), '跨类 i5 对 rc(i7) 最小编辑距离': int(min(cross.min(), cross2.min())), '同一序列与自身反向互补最小': int(own), '每周期全部达标': bool(ok),
            'GC范围': (min(sum(c in 'GC' for c in s) for s in S), max(sum(c in 'GC' for c in s) for s in S)), '同聚物个数': sum(bool(re.search(r'(.)\1\1', s)) for s in S),
            '序列数': len(set(S)), '与 Illumina 384 对相同序列数': sum(s in ill or rc(s) in ill for s in S)}
ver = {'96 组': verify(list(range(96))), '核心 48 组': verify(list(range(48)))}
f = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k}); hf = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
fl = {'A': PatternFill('solid', fgColor='E2EFDA'), 'B': PatternFill('solid', fgColor='FFF2CC'), 'C': PatternFill('solid', fgColor='F8CBAD')}
wb = Workbook(); s = wb.active; s.title = '说明'
tc96 = collections.Counter(info(k)['tier'] for k in range(96)); tc48 = collections.Counter(info(k)['tier'] for k in range(48)); srcs = collections.Counter(info(k)['src'] for k in range(96))
L = [('自建 UDP：重新配对后的 96 对（前 48 对 = 核心 48 组）', 1),
('做法：960 条自建序列（U001–U960，每条都可当 i5 或 i7）→ 第一阶段单个序列评估 → 第二阶段选 144 个 i5 + 144 个 i7 的候选池（编辑距离 ≥4、跨类反向互补 ≥3、每周期配色达标）→ 第三阶段对 100×100 子集算 i5×i7 的热点正向×反向相互作用 → 第四阶段在子集里选 96 对并配对（前 48 对为核心 48 组）。', 0),
('分级来源：%s；A：全局二聚体 ≥ −7 且 3′端 ≥ −4.5 kcal/mol；B：≥ −8.5 且 ≥ −5.5；C：其余（经验阈值，无尾巴基线 −5.7 / −3.8）。热点代理的预测与整管评估的一致率在 480 个固定配对管上为 99.6%%。' % '、'.join(f'{k} {v} 个管' for k, v in srcs.items()), 0),
('分级结果：96 组 A %d / B %d / C %d；核心 48 组 A %d / B %d / C %d。' % (tc96['A'], tc96['B'], tc96['C'], tc48['A'], tc48['B'], tc48['C']), 0)]
for k, v in ver.items(): L.append((f'{k}独立验证：' + '；'.join(f'{a}={b}' for a, b in v.items()), 0))
L.append(('命名：第 n 对的 i5 = DN5{n:03d}（放正向引物 5′端），i7 = DN7{n:03d}（放反向引物 5′端）；“U 编号”是 960 条序列的唯一编号（U001–U480 来自原 i5，U481–U960 来自原 i7，角色可以互换）。', 0))
L.append(('没有评估的：真实接头；实验验证（建议同样品多条形码对照）。逻辑和代码见 docs/primer_design/udp_denovo960/ 与 python_5R/explore/udp960_*.py。', 0))
s.column_dimensions['A'].width = 150
for i, (t, b) in enumerate(L, 1):
    c = s.cell(i, 1, t); c.font = f(bold=bool(b), size=12 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
w = wb.create_sheet('96对')
heads = ['序号', '核心48', 'i5 名称', 'i7 名称', 'i5 U编号', 'i7 U编号', 'i5 序列', 'i7 序列', 'i5 GC', 'i7 GC', '分级', '分级来源', '全局二聚体最差 ΔG', '3′端最差 ΔG', '发夹最差 ΔG', 'Olivar 总坏度']
for q in range(10): heads.append(f'i5 第{q + 1}位')
for q in range(10): heads.append(f'i7 第{q + 1}位')
for j, h in enumerate(heads, 1):
    c = w.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(wrap_text=True, horizontal='center'); c.border = bd
for k, p in enumerate(pairs):
    i = k + 2; m = info(k)
    vals = [k + 1, '是' if k < 48 else '', f'DN5{k + 1:03d}', f'DN7{k + 1:03d}', p['i5'], p['i7'], p['seq5'], p['seq7'], None, None, m['tier'], m['src'], round(m['H'], 2), round(m['E'], 2), None if m['hp'] is None else round(m['hp'], 2), None if m['bad'] is None else round(m['bad'])]
    for j, v in enumerate(vals, 1):
        c = w.cell(i, j, v); c.font = f(color='0000FF' if j in (7, 8) else '000000'); c.border = bd
    w.cell(i, 9, f'=LEN(SUBSTITUTE(SUBSTITUTE(G{i},"A",""),"T",""))'); w.cell(i, 10, f'=LEN(SUBSTITUTE(SUBSTITUTE(H{i},"A",""),"T",""))'); w.cell(i, 11).fill = fl[m['tier']]
    for q in range(10): w.cell(i, 17 + q, f'=MID($G{i},{q + 1},1)'); w.cell(i, 27 + q, f'=MID($H{i},{q + 1},1)')
for j, wd in enumerate([6, 7, 9, 9, 8, 8, 14, 14, 6, 6, 6, 8, 11, 11, 10, 11], 1): w.column_dimensions[get_column_letter(j)].width = wd
w.row_dimensions[1].height = 40; w.freeze_panes = 'G2'
c1 = wb.create_sheet('每周期统计')
for j, h in enumerate(['范围', '序列', '周期', 'A', 'C', 'G', 'T', 'GC%', 'A+C%', 'A+T%', '是否达标'], 1):
    c = c1.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf
row = 2
for lab, N in (('96 组', 96), ('核心 48 组', 48)):
    for name, base in (('i5', 17), ('i7', 27)):
        for p in range(10):
            hc = get_column_letter(base + p); rng = f"'96对'!${hc}$2:${hc}${N + 1}"
            c1.cell(row, 1, lab); c1.cell(row, 2, name); c1.cell(row, 3, p + 1)
            for j, b in enumerate('ACGT', 4): c1.cell(row, j, f'=COUNTIF({rng},"{b}")')
            c1.cell(row, 8, f'=(E{row}+F{row})/{N}'); c1.cell(row, 9, f'=(D{row}+E{row})/{N}'); c1.cell(row, 10, f'=(D{row}+G{row})/{N}')
            c1.cell(row, 11, f'=IF(AND(H{row}>=0.4,H{row}<=0.6,I{row}>=0.4,I{row}<=0.6,J{row}>=0.4,J{row}<=0.6,MIN(D{row}:G{row})>={round(0.15 * N, 2)}),"是","否")')
            for j in (8, 9, 10): c1.cell(row, j).number_format = '0.0%'
            row += 1
for r_ in c1.iter_rows(min_row=2):
    for c in r_: c.font = f()
wb.save(outx)
ws = openpyxl.load_workbook(D + 'primer_order_6amp_pool33_v3.xlsx')['订购清单']; oligos = [(r_[1], r_[2], r_[4]) for r_ in ws.iter_rows(min_row=4, max_row=36, values_only=True)]
with open(outc, 'w', newline='', encoding='utf-8-sig') as fh:
    cw = csv.writer(fh); cw.writerow(['对', '核心48', 'i5 名称', 'i7 名称', '分级', '引物名', '位点', '尾巴', '特异部分', '完整序列 5′→3′', '长度'])
    for k, p in enumerate(pairs):
        for nm, site, q in oligos:
            t = p['seq5'] if site.endswith('F') else p['seq7']; cw.writerow([k + 1, '是' if k < 48 else '', f'DN5{k + 1:03d}', f'DN7{k + 1:03d}', info(k)['tier'], f'{nm}-{"DN5" if site.endswith("F") else "DN7"}{k + 1:03d}', site, t, q, t + q, len(t + q)])
print(json.dumps(ver, ensure_ascii=False, indent=1)); print('96 组', dict(tc96), '核心 48', dict(tc48), dict(srcs))
