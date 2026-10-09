import sys, json, pickle, csv
import numpy as np, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SC = sys.argv[1]; D = '/home/user/16S_5R/docs/primer_design/'
V = json.load(open(SC + '/final_verify.json')); sel48, sel96 = V['sel48'], V['sel96']; T = V['tubes']
U = {l.split()[0]: l.split() for l in open(D + 'udp_tail/udp_all_1_384.tsv')}
off = {}
for tag in ('human', 'pig'):
    o = pickle.load(open(f'{SC}/off96_{tag}.pkl', 'rb')); names = [l.split()[0] for l in open(D + 'udp_tail/udp_selected_96.tsv')]
    g = np.array([r[2] for r in o['rows']]); tot = g[:, None] + o['TM'].astype(int)
    for k, n in enumerate(names): off[(tag, n)] = (int((tot[:, k] <= 3).sum()), int((tot[:, k] <= 4).sum()))
f = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k}); hf = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
fl = {'A': PatternFill('solid', fgColor='E2EFDA'), 'B': PatternFill('solid', fgColor='FFF2CC')}
wb = Workbook(); s = wb.active; s.title = '说明'
o48, o96 = V['out']['48'], V['out']['96']
L = [('Illumina UDP（Set A–D）48 组与 96 组内联索引推荐', 1),
('来源：Illumina Adapter Sequences（文档 1000000002694 v22）的 UDP0001–0384，i5/i7 取 “Bases in Adapter” 列，i5 加在正向引物 5′端，同编号 i7 加在反向引物 5′端，没有接头。已排除 UDP0017（你们测试测序量显著降低）。Oligonucleotide sequences © 2025 Illumina, Inc. All rights reserved.', 0),
('筛选与硬约束：① i5 之间、i7 之间、i5 与 i7 之间的 Levenshtein 编辑距离都 ≥4（实测最小值 4）；② 每个周期（第 1–10 位）GC 占比、A+C 占比、A+T 占比都在 40–60%，A/C/G/T 每种 ≥15%（i5、i7 各自满足）；③ 近回文索引已剔除，无同聚物 ≥3；④ 每个管（同编号 i5+i7 加到 33 条引物池）的二聚体分级为 A 或 B，没有 C 级。', 0),
('分级（经验阈值，没有文献依据）：A = 管内全局二聚体 ΔG ≥ −7 且 3′端 ≥ −4.5（58 °C，primer3）；B = ≥ −8.5 且 ≥ −5.5；无尾巴基线为 −5.7 / −3.8。', 0),
('结果：48 组 = 28 个 A + 20 个 B；96 组 = 40 个 A + 56 个 B，48 组是 96 组的子集（先做 48 再扩到 96 即可）。', 0),
('为什么 48 组不能只含 A 级：A 级候选的 i5 第 10 位（紧邻引物）只有 A/T，这个周期 GC 为 0%，C/G 没有信号；i5 以 G/C 结尾的候选全是 B 级，满足 GC ≥40% 至少要 20 个，所以 20 个 B 是下限。', 0),
('提醒：① 部分索引单个 GC 为 3 或 7（占 10 位的 30% 或 70%）：48 组 i5 17 个、i7 16 个，96 组 i5 29 个、i7 31 个，每个周期仍达标；② i5 与 rc(i7) 的最小编辑距离只有 2，ONT 反向读取时要靠引物序列判断方向；③ 二聚体、发夹按 33 条引物池 + 单个管内的 i5/i7 评估，没有真实接头；④ 脱靶：尾巴放在 5′端，不改变 3′端种子，人和猪基因组全长（含尾巴）≤3 错配的位点见 96 组表；⑤ 混合建库后样本间 UDP 互补：头尾配对最长互补 7 nt、60 °C 最差 ΔG −5.4，锅柄最差 −3.1，没有 10 nt 完全互补（原始 UDP0001–0096 有 16 对）。', 0)]
s.column_dimensions['A'].width = 150
for i, (t, b) in enumerate(L, 1):
    c = s.cell(i, 1, t); c.font = f(bold=bool(b), size=12 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
w = wb.create_sheet('96组')
H = ['序号', 'UDP', '在48组', 'i5（正向引物尾巴）', 'i7（反向引物尾巴）', 'i5 GC', 'i7 GC', '分级', '管内全局二聚体最差 ΔG', '3′端最差 ΔG', '发夹最差 ΔG', 'Olivar 总坏度', '人 全长≤3错配位点', '人 ≤4', '猪 全长≤3错配位点', '猪 ≤4']
for j, h in enumerate(H, 1):
    c = w.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(wrap_text=True, horizontal='center'); c.border = bd
for i, n in enumerate(sorted(sel96, key=lambda x: (x not in sel48, int(x[3:7]))), 2):
    t = T[n]; r = [i - 1, n, '是' if n in sel48 else '', U[n][3], U[n][1], None, None, t['tier'], round(t['minH'], 2), round(t['minE'], 2), round(t['hp'], 2), round(t['bad']), *off[('human', n)], *off[('pig', n)]]
    for j, v in enumerate(r, 1):
        c = w.cell(i, j, v); c.font = f(color='0000FF' if j in (4, 5) else '000000'); c.border = bd
    w.cell(i, 6, f'=LEN(SUBSTITUTE(SUBSTITUTE(D{i},"A",""),"T",""))'); w.cell(i, 7, f'=LEN(SUBSTITUTE(SUBSTITUTE(E{i},"A",""),"T",""))')
    w.cell(i, 8).fill = fl[t['tier']]
    for q in range(10):
        w.cell(i, 17 + q, f'=MID($D{i},{q + 1},1)'); w.cell(i, 27 + q, f'=MID($E{i},{q + 1},1)')
for q in range(10):
    w.cell(1, 17 + q, f'i5 第{q + 1}位'); w.cell(1, 27 + q, f'i7 第{q + 1}位')
for j, wd in enumerate([6, 12, 8, 14, 14, 6, 6, 6, 12, 10, 10, 12, 10, 8, 10, 8], 1): w.column_dimensions[get_column_letter(j)].width = wd
w.row_dimensions[1].height = 45; w.freeze_panes = 'D2'
# 每周期（公式）
c1 = wb.create_sheet('每周期统计')
heads = ['集合', '序列', '周期', 'A', 'C', 'G', 'T', 'GC%', 'A+C%', 'A+T%', '是否达标']
for j, h in enumerate(heads, 1):
    c = c1.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf
row = 2
for lab, rng_end, flt in (('48', 49, '是'), ('96', 97, None)):
    for col, name in (('D', 'i5'), ('E', 'i7')):
        for p in range(1, 11):
            hc = get_column_letter((17 if name == 'i5' else 27) + p - 1); rngs = f"'96组'!${hc}$2:${hc}${rng_end}"; N = 48 if lab == '48' else 96
            c1.cell(row, 1, lab + ' 组'); c1.cell(row, 2, name); c1.cell(row, 3, p)
            for j, b in enumerate('ACGT', 4): c1.cell(row, j, f'=COUNTIF({rngs},"{b}")')
            c1.cell(row, 8, f'=(E{row}+F{row})/{N}'); c1.cell(row, 9, f'=(D{row}+E{row})/{N}'); c1.cell(row, 10, f'=(D{row}+G{row})/{N}')
            c1.cell(row, 11, f'=IF(AND(H{row}>=0.4,H{row}<=0.6,I{row}>=0.4,I{row}<=0.6,J{row}>=0.4,J{row}<=0.6,MIN(D{row}:G{row})>={round(0.15 * N, 2)}),"是","否")')
            for j in (8, 9, 10): c1.cell(row, j).number_format = '0.0%'
            row += 1
for r in c1.iter_rows(min_row=2):
    for c in r: c.font = f()
c1.cell(row + 1, 1, '48 组位于 96 组表的前 48 行（“在48组”=是的排在最前）；公式统计的是对应行范围。').font = f()
# 完整引物 CSV（96 管 × 33 条）
ws = openpyxl.load_workbook(D + 'primer_order_6amp_pool33_v3.xlsx')['订购清单']
oligos = [(r[1], r[2], r[4]) for r in ws.iter_rows(min_row=4, max_row=36, values_only=True)]
with open(D + 'udp_final_full_oligos_96.csv', 'w', newline='', encoding='utf-8-sig') as fh:
    cw = csv.writer(fh); cw.writerow(['UDP', '在48组', '引物名', '位点', '尾巴', '特异部分', '完整序列 5′→3′', '长度'])
    for n in sorted(sel96, key=lambda x: (x not in sel48, int(x[3:7]))):
        for nm, site, q in oligos:
            t = U[n][3] if site.endswith('F') else U[n][1]; cw.writerow([n, '是' if n in sel48 else '', f'{nm}-{n}', site, t, q, t + q, len(t + q)])
wb.save(D + 'udp_final_48_96_selection.xlsx'); print('ok')
