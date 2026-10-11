import sys, pickle, json, random, math, csv
import numpy as np, pandas as pd, openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SC = sys.argv[1]; D = '/home/user/16S_5R/docs/primer_design/'
X = pd.read_pickle(SC + '/dn_b1_metrics.pkl').sort_values('name').reset_index(drop=True)
OFF = pickle.load(open(SC + '/off_dn_sum.pkl', 'rb'))
X['k'] = X.name.str[2:].astype(int); X['n5'] = X.k.map(lambda k: f'DN5{k:02d}'); X['n7'] = X.k.map(lambda k: f'DN7{k:02d}')
B = {'A': 0, 'C': 1, 'G': 2, 'T': 3}; M5 = np.array([[B[c] for c in s] for s in X.i5]); M7 = np.array([[B[c] for c in s] for s in X.i7])
def pen(S, N):
    p = 0.0
    for M in (M5, M7):
        for pos in range(10):
            c = np.bincount(M[S, pos], minlength=4); gc = c[1] + c[2]; red = c[0] + c[1]; green = c[0] + c[3]
            for v in (gc, red, green): p += max(0, 0.4 * N - v) ** 2 + max(0, v - 0.6 * N) ** 2
            for v in c: p += max(0, 0.15 * N - v) ** 2
    return p
ok = X.index.tolist(); isA = (X.tier == 'A').values; cst = X.tier.map({'A': 0, 'B': 3, 'C': 30}).values
def best_subset(N, pool, seeds=8, iters=40000):
    best = None
    for sd in range(seeds):
        rnd = random.Random(sd); S = rnd.sample(pool, N); rest = [x for x in pool if x not in S]
        e = 1000 * pen(S, N) + cst[S].sum()
        for it in range(iters):
            T = 30 * (1 - it / iters) + 0.05; i = rnd.randrange(N); j = rnd.randrange(len(rest)); a = S[i]; S[i] = rest[j]
            e2 = 1000 * pen(S, N) + cst[S].sum()
            if e2 <= e or rnd.random() < math.exp((e - e2) / T): rest[j] = a; e = e2
            else: S[i] = a
            if best is None or e < best[0]: best = (e, list(S))
    return best
b48 = best_subset(48, ok); S48 = b48[1]; nC = int((X.tier.values[S48] == 'C').sum()); print('48 组：越界', pen(S48, 48), 'A', int(isA[S48].sum()), 'C', nC)
X['rec48'] = ''; X.loc[S48, 'rec48'] = '推荐'
f = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k}); hf = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
fl = {'A': PatternFill('solid', fgColor='E2EFDA'), 'B': PatternFill('solid', fgColor='FFF2CC'), 'C': PatternFill('solid', fgColor='F8CBAD')}
wb = Workbook(); s = wb.active; s.title = '说明'
A_, B_, C_ = [int((X.tier == t).sum()) for t in 'ABC']
L = [('自建 UDP（DN）第 1 块 96 对（DN501–DN596 放正向引物，DN701–DN796 放反向引物）的管内评估', 1),
('做法：同编号的 DN5xx（i5）加在池中 60 个正向展开序列的 5′端，DN7xx（i7）加在 138 个反向展开序列的 5′端，198 个展开序列两两评估（primer3，58 °C，50 mM Na⁺ / 2 mM Mg²⁺ / 0.2 mM dNTP / 250 nM），再与无尾巴基线比较；没有真实接头。', 0),
('基线（无尾巴）：全局二聚体最差 ΔG −5.7，3′端锚定 −3.8，发夹 −1.2，Olivar 总坏度 35652。', 0),
('分级（经验阈值，没有文献依据）：A = 全局 ≥ −7 且 3′端 ≥ −4.5；B = ≥ −8.5 且 ≥ −5.5；C = 其余。结果：A %d、B %d、C %d。' % (A_, B_, C_), 0),
('重要更正：我之前说自建序列可以做到每个管都是 A 级，这个判断没有被评估支持：自建序列的 A 级比例（%d/96 = %d%%）和 Illumina 序列（86/371 = 23%%）差不多。预筛指标（尾巴对池里引物的二聚体）不能预测管内结果；最差的相互作用多数是正向×反向（%d 个管）。' % (A_, round(100 * A_ / 96), (X[['H_FF', 'H_RR', 'H_FR']].idxmin(axis=1) == 'H_FR').sum()), 0),
('推荐 48 组：在本块 96 个管里挑 48 个，要求每个周期 GC、A+C、A+T 占比 40–60%%、A/C/G/T 各 ≥15%%（i5、i7 各自）且尽量少用差的管：含 %d 个 A 级、%d 个 B 级、%d 个 C 级；周期越界量 %.1f（0 表示全部达标）。仅用 A/B 级时本块凑不出满足每周期要求的 48 个。' % (int(isA[S48].sum()), 48 - int(isA[S48].sum()) - nC, nC, pen(S48, 48)), 0),
('脱靶：尾巴在 5′端，3′端种子不变；把尾巴并入后，人/猪基因组全长 ≤2 错配位点在所有管里都是 0，≤3 错配最多 人 %d / 猪 %d 个。' % (OFF['human'][0].max(), OFF['pig'][0].max()), 0),
('混合建库后的 UDP 互补（96 对）：头尾配对最长互补 8 nt，60 °C 最差 ΔG −4.5；锅柄最差 −2.4；没有 10 nt 完全互补。', 0),
('下一步建议：重新配对（F×R 的相互作用是配对相关的）并从 384 个里替换 C 级管；我可以按“每个 i5 单独评估正向×正向、每个 i7 单独评估反向×反向、再优化配对”做，但需要几小时的计算，请你确认。', 0)]
s.column_dimensions['A'].width = 150
for i, (t, b) in enumerate(L, 1):
    c = s.cell(i, 1, t); c.font = f(bold=bool(b), size=12 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
w = wb.create_sheet('96管评估')
heads = ['i5 名称', 'i7 名称', 'i5 序列', 'i7 序列', '分级', '推荐48', '全局二聚体最差 ΔG', '3′端最差 ΔG', '正向×正向', '反向×反向', '正向×反向', '发夹最差 ΔG', '发夹≤−3 的展开数', '尾巴×尾巴 ΔG', 'Olivar 总坏度', '人 全长≤3错配', '人 ≤4', '猪 全长≤3错配', '猪 ≤4']
for j, h in enumerate(heads, 1):
    c = w.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(wrap_text=True, horizontal='center'); c.border = bd
for i, r in X.iterrows():
    vals = [r.n5, r.n7, r.i5, r.i7, r.tier, r.rec48, round(r.minH, 2), round(r.minE, 2), round(r.H_FF, 2), round(r.H_RR, 2), round(r.H_FR, 2), round(r.hp, 2), int(r.nhp3), round(r.tt, 2), round(r.bad),
            int(OFF['human'][0][i]), int(OFF['human'][1][i]), int(OFF['pig'][0][i]), int(OFF['pig'][1][i])]
    for j, v in enumerate(vals, 1):
        c = w.cell(i + 2, j, v); c.font = f(color='0000FF' if j in (3, 4) else '000000'); c.border = bd
    w.cell(i + 2, 5).fill = fl[r.tier]
for j, wd in enumerate([9, 9, 14, 14, 6, 8, 10, 10, 9, 9, 9, 9, 9, 9, 11, 9, 7, 9, 7], 1): w.column_dimensions[get_column_letter(j)].width = wd
w.row_dimensions[1].height = 45; w.freeze_panes = 'C2'
st = wb.create_sheet('统计'); st['A1'] = '项目'; st['B1'] = '数量'
for c in ('A1', 'B1'): st[c].font = f(bold=True)
for i, (a, b) in enumerate([('A 级', "=COUNTIF('96管评估'!E2:E97,\"A\")"), ('B 级', "=COUNTIF('96管评估'!E2:E97,\"B\")"), ('C 级', "=COUNTIF('96管评估'!E2:E97,\"C\")"), ('推荐 48 组', "=COUNTIF('96管评估'!F2:F97,\"推荐\")"),
                            ('推荐中 A 级', "=COUNTIFS('96管评估'!F2:F97,\"推荐\",'96管评估'!E2:E97,\"A\")"), ('全局 ΔG ≤ −8 的管', "=COUNTIF('96管评估'!G2:G97,\"<=-8\")"), ('3′ ΔG ≤ −5 的管', "=COUNTIF('96管评估'!H2:H97,\"<=-5\")")], 2):
    st.cell(i, 1, a).font = f(); st.cell(i, 2, b).font = f()
st.column_dimensions['A'].width = 24
wb.save(D + 'udp_denovo_block1_evaluation.xlsx')
ws = openpyxl.load_workbook(D + 'primer_order_6amp_pool33_v3.xlsx')['订购清单']; oligos = [(r[1], r[2], r[4]) for r in ws.iter_rows(min_row=4, max_row=36, values_only=True)]
with open(D + 'udp_denovo_block1_full_oligos.csv', 'w', newline='', encoding='utf-8-sig') as fh:
    cw = csv.writer(fh); cw.writerow(['管', 'i5 名称', 'i7 名称', '分级', '推荐48', '引物名', '位点', '尾巴', '特异部分', '完整序列 5′→3′', '长度'])
    for i, r in X.iterrows():
        for nm, site, q in oligos:
            t = r.i5 if site.endswith('F') else r.i7; cw.writerow([r['name'], r.n5, r.n7, r.tier, r.rec48, f'{nm}-{r.n5 if site.endswith("F") else r.n7}', site, t, q, t + q, len(t + q)])
print('ok', A_, B_, C_)
