"""把选出的 48 / 96 组导出为 Excel（含独立验证、每周期公式统计）和完整引物 csv。
用法：python3 udp_dn_export.py <dn480_final.json> <out xlsx> <out csv> <sel48 json|-> <sel96 json|-> <off dir 前缀|-> <pkl...>
 sel json 来自 udp_dn_select_fast.py（res[N].idx = 管编号-1）。命名：管编号 k（DN001–DN480）对应 i5 = DN5{k:03d}、i7 = DN7{k:03d}。"""
import sys, json, pickle, collections, re, csv
import numpy as np, openpyxl
from rapidfuzz.process import cdist
from rapidfuzz.distance import Levenshtein
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
fin, outx, outc, s48, s96, offp = sys.argv[1:7]; pkls = sys.argv[7:]
D = '/home/user/16S_5R/docs/primer_design/'
d = json.load(open(fin)); i5a, i7a = d['i5'], d['i7']; rc = lambda s: s[::-1].translate(str.maketrans('ACGT', 'TGCA'))
tubes = {}
for p in pkls:
    r = pickle.load(open(p, 'rb')); N0 = len(r['EX'])
    for t in r['tubes'].values(): tubes[int(t['name'][2:])] = t
iu = np.triu_indices(N0); iF = np.array([x[1] == 'F' for x in r['EX']])
def metrics(k):
    t = tubes[k]; h = float(t['H'][iu].min()); e = float(t['E'].min())
    return dict(tier='A' if (h >= -7 and e >= -4.5) else 'B' if (h >= -8.5 and e >= -5.5) else 'C', minH=h, minE=e, hp=float(t['hg'].min()), bad=float(t['bad']))
sets = {}
for lab, path in (('48', s48), ('96', s96)):
    if path != '-': sets[lab] = sorted(k + 1 for k in json.load(open(path))['res'][lab]['idx'])
def verify(ks):
    i5 = [i5a[k - 1] for k in ks]; i7 = [i7a[k - 1] for k in ks]; S = i5 + i7; n = len(S); N = len(ks)
    Dm = cdist(S, S, scorer=Levenshtein.distance, dtype=np.uint8, workers=2); np.fill_diagonal(Dm, 99)
    Rm = cdist(S, [rc(s) for s in S], scorer=Levenshtein.distance, dtype=np.uint8, workers=2); np.fill_diagonal(Rm, 99)
    own = [Levenshtein.distance(s, rc(s)) for s in S]
    ok = True
    for M in (i5, i7):
        for p in range(10):
            c = collections.Counter(s[p] for s in M); v = (c['G'] + c['C'], c['A'] + c['C'], c['A'] + c['T'])
            ok &= all(0.4 * N <= x <= 0.6 * N for x in v) and min(c[b] for b in 'ACGT') >= 0.15 * N
    return dict(最小编辑距离=int(Dm.min()), 最小反向互补距离=int(min(Rm.min(), min(own))), 每周期全部达标=bool(ok), GC范围=(min(sum(c in 'GC' for c in s) for s in S), max(sum(c in 'GC' for c in s) for s in S)),
                同聚物个数=sum(bool(re.search(r'(.)\1\1', s)) for s in S), 序列数=len(set(S)))
f = lambda **k: Font(**{'name': 'Arial', 'size': 10, **k}); hf = PatternFill('solid', fgColor='1F3864')
thin = Side(style='thin', color='BFBFBF'); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
fl = {'A': PatternFill('solid', fgColor='E2EFDA'), 'B': PatternFill('solid', fgColor='FFF2CC'), 'C': PatternFill('solid', fgColor='F8CBAD')}
wb = Workbook(); s = wb.active; s.title = '说明'
ver = {lab: verify(ks) for lab, ks in sets.items()}
L = [('自建 UDP（DN）从 480 对候选中选出的 48 / 96 组', 1),
('命名：候选管编号 k（DN001–DN480）对应 i5 = DN5{k:03d}（放正向引物 5′端）、i7 = DN7{k:03d}（放反向引物 5′端），例如 k=1 → DN5001 / DN7001。i5、i7 按候选管配对，只用同一管里的一对做管内评估。', 0),
('选择逻辑：① 480 个候选管都按 33 条引物池（60 个正向、138 个反向展开序列）加尾巴后做整管评估（primer3，58 °C）；② 分级 A：全局二聚体 ΔG ≥ −7 且 3′端 ≥ −4.5；B：≥ −8.5 且 ≥ −5.5；C：其余（阈值是经验阈值，基线无尾巴为 −5.7 / −3.8）；③ 在 A/B 级管里用模拟退火选 N 个管，硬约束：所选 i5∪i7 两两 Levenshtein ≥4、与对方反向互补 ≥4；每个周期（第 1–10 位）GC、A+C、A+T 占比 40–60%、A/C/G/T 各 ≥15%（i5、i7 各自）；软目标：A 级尽量多、二聚体尽量弱。', 0),
('候选管总分级：A 94、B 171、C 215（480 个）。A 级总数不够，且全 A 级的子集无法满足每周期配色和距离约束，所以 48 / 96 组里都需要 B 级管。', 0)]
for lab, v in ver.items():
    tc = collections.Counter(metrics(k)['tier'] for k in sets[lab]); L.append((f'{lab} 组独立验证：' + '；'.join(f'{a}={b}' for a, b in v.items()) + f'；分级 A {tc["A"]} / B {tc["B"]} / C {tc["C"]}', 0))
L.append(('没有评估的：真实接头；实验验证（条形码紧邻引物会带来扩增偏好，建议同样品多条形码做对照）。代码见 python_5R/explore/udp_dn_select_fast.py、udp_dn_export.py；逻辑说明见 docs/primer_design/udp_denovo/SELECTION.md。', 0))
s.column_dimensions['A'].width = 150
for i, (t, b) in enumerate(L, 1):
    c = s.cell(i, 1, t); c.font = f(bold=bool(b), size=12 if b else 10); c.alignment = Alignment(wrap_text=True, vertical='top')
offs = {}
if offp != '-':
    for tag in ('human', 'pig'):
        try:
            o = pickle.load(open(f'{offp}_{tag}.pkl', 'rb')); g = np.array([r_[2] for r_ in o['rows']]); tot = g[:, None] + o['TM'].astype(int)
            names = [l.split()[0] for l in open(D + 'udp_tail/udp_denovo_sel48.tsv')]
            offs[tag] = {int(nm[2:]): (int((tot[:, j] <= 3).sum()), int((tot[:, j] <= 4).sum())) for j, nm in enumerate(names)}
        except Exception as e: pass
for lab, ks in sets.items():
    w = wb.create_sheet(f'{lab}组'); heads = ['序号', '管编号', 'i5 名称', 'i7 名称', 'i5 序列', 'i7 序列', 'i5 GC', 'i7 GC', '分级', '全局二聚体最差 ΔG', '3′端最差 ΔG', '发夹最差 ΔG', 'Olivar 总坏度', '人 全长≤3错配', '猪 全长≤3错配']
    for q in range(10): heads.append(f'i5 第{q + 1}位')
    for q in range(10): heads.append(f'i7 第{q + 1}位')
    for j, h in enumerate(heads, 1):
        c = w.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(wrap_text=True, horizontal='center'); c.border = bd
    for i, k in enumerate(ks, 2):
        m = metrics(k); vals = [i - 1, f'DN{k:03d}', f'DN5{k:03d}', f'DN7{k:03d}', i5a[k - 1], i7a[k - 1], None, None, m['tier'], round(m['minH'], 2), round(m['minE'], 2), round(m['hp'], 2), round(m['bad']),
                                offs.get('human', {}).get(k, ('', ''))[0], offs.get('pig', {}).get(k, ('', ''))[0]]
        for j, v in enumerate(vals, 1):
            c = w.cell(i, j, v); c.font = f(color='0000FF' if j in (5, 6) else '000000'); c.border = bd
        w.cell(i, 9).fill = fl[m['tier']]
        w.cell(i, 7, f'=LEN(SUBSTITUTE(SUBSTITUTE(E{i},"A",""),"T",""))'); w.cell(i, 8, f'=LEN(SUBSTITUTE(SUBSTITUTE(F{i},"A",""),"T",""))')
        for q in range(10): w.cell(i, 16 + q, f'=MID($E{i},{q + 1},1)'); w.cell(i, 26 + q, f'=MID($F{i},{q + 1},1)')
    for j, wd in enumerate([6, 8, 9, 9, 14, 14, 6, 6, 6, 10, 10, 10, 11, 9, 9], 1): w.column_dimensions[get_column_letter(j)].width = wd
    w.row_dimensions[1].height = 40; w.freeze_panes = 'E2'
    c1 = wb.create_sheet(f'{lab}组每周期'); N = len(ks)
    for j, h in enumerate(['序列', '周期', 'A', 'C', 'G', 'T', 'GC%', 'A+C%', 'A+T%', '是否达标'], 1):
        c = c1.cell(1, j, h); c.font = f(bold=True, color='FFFFFF'); c.fill = hf
    row = 2
    for name, base in (('i5', 16), ('i7', 26)):
        for p in range(10):
            hc = get_column_letter(base + p); rng = f"'{lab}组'!${hc}$2:${hc}${N + 1}"
            c1.cell(row, 1, name); c1.cell(row, 2, p + 1)
            for j, b in enumerate('ACGT', 3): c1.cell(row, j, f'=COUNTIF({rng},"{b}")')
            c1.cell(row, 7, f'=(D{row}+E{row})/{N}'); c1.cell(row, 8, f'=(C{row}+D{row})/{N}'); c1.cell(row, 9, f'=(C{row}+F{row})/{N}')
            c1.cell(row, 10, f'=IF(AND(G{row}>=0.4,G{row}<=0.6,H{row}>=0.4,H{row}<=0.6,I{row}>=0.4,I{row}<=0.6,MIN(C{row}:F{row})>={round(0.15 * N, 2)}),"是","否")')
            for j in (7, 8, 9): c1.cell(row, j).number_format = '0.0%'
            row += 1
    for r_ in c1.iter_rows(min_row=2):
        for c in r_: c.font = f()
wb.save(outx)
ws = openpyxl.load_workbook(D + 'primer_order_6amp_pool33_v3.xlsx')['订购清单']; oligos = [(r_[1], r_[2], r_[4]) for r_ in ws.iter_rows(min_row=4, max_row=36, values_only=True)]
with open(outc, 'w', newline='', encoding='utf-8-sig') as fh:
    cw = csv.writer(fh); cw.writerow(['组', '管编号', 'i5 名称', 'i7 名称', '分级', '引物名', '位点', '尾巴', '特异部分', '完整序列 5′→3′', '长度'])
    for lab, ks in sets.items():
        for k in ks:
            m = metrics(k)
            for nm, site, q in oligos:
                t = i5a[k - 1] if site.endswith('F') else i7a[k - 1]; cw.writerow([lab, f'DN{k:03d}', f'DN5{k:03d}', f'DN7{k:03d}', m['tier'], f'{nm}-{"DN5" if site.endswith("F") else "DN7"}{k:03d}', site, t, q, t + q, len(t + q)])
print(json.dumps(ver, ensure_ascii=False, indent=1))
