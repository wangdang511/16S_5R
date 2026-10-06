# 生成 docs/primer_design/primer_order_12_16_24.xlsx（openpyxl；公式用独立求值器核对过）
import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
D='/home/user/16S_5R/docs/primer_design/'
W=pd.read_csv(D+'order_list_levels.csv'); Lg=pd.read_csv(D+'order_list_levels_long.csv'); Lg['level']=Lg.level.astype(str)
sh={lv:Lg[Lg.level==lv].set_index('order').share for lv in ('24','16','12')}
REG={'A1':'V1·V2','A2':'V3','A3':'V4','V5':'V5','A4':'V6·V7','A5':'V8·V9'}
def ascii_name(o):
    n=o.replace('SMURF5-','S5-').replace('（SNAP）','').replace('（补充）','s').replace('A1-F.V1f','A1-F.V1f').replace('V5-R.907R','V5-R.907R')
    return n
W['name_ascii']=W.order.map(ascii_name)
W['region']=W.slot.map(lambda s:REG[s.split('-')[0]]); W['orient']=W.slot.map(lambda s:'正向 F' if s.endswith('F') else '反向 R')
W['type']=np.where(W.is_sup,'支原体补充','核心/借鉴')
W.loc[W.order.str.contains('SNAP'),'type']='借鉴 Swift SNAP'
FONT='Arial'
f=lambda **k: Font(name=FONT,**k)
hdr_fill=PatternFill('solid',fgColor='1F3864'); hdr_font=f(bold=True,color='FFFFFF',size=10)
inp_font=f(color='0000FF',size=10); base=f(size=10); bold=f(bold=True,size=10)
yellow=PatternFill('solid',fgColor='FFF2CC'); grey=PatternFill('solid',fgColor='F2F2F2')
thin=Side(style='thin',color='BFBFBF'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical='top'); center=Alignment(horizontal='center',vertical='center',wrap_text=True)
wb=Workbook()
# ---------- 说明
ws=wb.active; ws.title='说明'
lines=[('16S 引物订购清单：12、16、24 条引物池',True),
('用途：一次订齐 24 条寡核苷酸；12、16、24 条三个池都是这 24 条的子集，按各池工作表里的体积配成 100 µL 的 10× 混合液。',False),
('',False),
('工作表',True),
('订购总表：24 条寡核苷酸（含序列、位置、Tm、展开数、各池是否包含），可直接作为订购单。',False),
('12 条池 / 16 条池 / 24 条池：各池的配制表（每条引物加多少 µL 的母液、位点内比例、总量核对、加水量）。',False),
('池比较：三个池的评估摘要（来自计算评估，非实测）。',False),
('',False),
('颜色约定',True),
('蓝色字 = 输入（可以修改）：合成规模、纯化方式、母液浓度、10× 混合液体积、各引物位点内比例。黑色 = 公式，自动计算。黄色底 = 订购前需要你确认的设置。',False),
('',False),
('订购注意',True),
('1. 序列用 IUPAC 简并码（R=A/G，Y=C/T，W=A/T，M=A/C，K=G/T，S=G/C，D=A/G/T，B=C/G/T，V=A/C/G），按“混合碱基”合成，每条订一个管，不要拆成单个展开序列。订购前请确认合成商支持这些简并碱基。',False),
('2. 序列只含基因特异部分，没有接头、UDP 或修饰；如果建库要在引物上加接头，请自己加，并重新检查 Tm 和二聚体（我没有评估带接头的版本）。',False),
('3. 合成规模和纯化方式不是我推荐的结果，是占位的默认值（25 nmol、标准脱盐），请按你的合成商和用量改成实际要求；含有较多简并碱基的引物，合成商可能建议更高的规模。',False),
('4. “订购名称”用 ASCII 字符（S5- 开头，末尾 s = 支原体补充），可以直接用于订购；“中文名”只是对应页面里的名称。',False),
('5. 配池：每个引物位点（共 11 个）总量相同；同位点内多条时按各自覆盖比例分，补充引物占 10%；24 条池里同位点的冗余引物（A1-F 与 V1_f；V5-R.1、V5-R.2 与 907R）按覆盖比例平分。',False),
('6. 这些设计是序列层面的计算结果（Greengenes 13_8 和 SILVA 128；匹配规则：最多 1 个错配、3′ 端 3 个碱基匹配），没有实验验证；退火温度建议先做梯度（大致 52–60 °C）。',False),
('',False),
('数据来源：docs/primer_design/order_list_levels.csv、order_list_levels_long.csv，评估见 primer_order_list_levels.html。',False)]
ws.column_dimensions['A'].width=130
for i,(t,b) in enumerate(lines,1):
    c=ws.cell(row=i,column=1,value=t); c.font=f(bold=b,size=12 if i==1 else 10); c.alignment=Alignment(wrap_text=True,vertical='top')
# ---------- 订购总表
wo=wb.create_sheet('订购总表')
heads=['序号','订购名称（ASCII）','中文名','序列 5′→3′（IUPAC）','长度 nt','引物位点','区域','方向','E. coli 位置','平均 Tm °C','Tm 范围 °C（展开序列）','GC 含量','展开序列数','类型','人基因组位点 / 展开序列（≤2 错配）','在 12 条池','在 16 条池','在 24 条池','合成规模','纯化方式','修饰','备注']
for j,h in enumerate(heads,1):
    c=wo.cell(row=1,column=j,value=h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=center; c.border=border
note={'A1-F':'与 V1_f 同位点，二选一（24 条池同时含两者）','V5-R.1':'','V5-R.2':'','A2-R':'V3 反向引物（长版，516–535，扩增子 338–535）'}
r0=2
for i,r in enumerate(W.itertuples(),r0):
    vals=[i-1,r.name_ascii,r.order,r.seq,None,r.slot,r.region,r.orient,r.pos,r.Tm,f'{r.Tm_lo}–{r.Tm_hi}',r.GC,int(r.expansions),r.type,float(r.human_per_exp),
          '✓' if pd.notna(r.uL_12) else '','✓' if pd.notna(r.uL_16) else '','✓' if pd.notna(r.uL_24) else '',None,None,'无','']
    for j,v in enumerate(vals,1):
        c=wo.cell(row=i,column=j,value=v); c.font=base; c.border=border; c.alignment=Alignment(vertical='top',wrap_text=(j in (3,22)))
    wo.cell(row=i,column=5,value=f'=LEN(D{i})').font=base
    wo.cell(row=i,column=19,value=25).font=inp_font; wo.cell(row=i,column=19).number_format='0 "nmol"'
    wo.cell(row=i,column=20,value='标准脱盐（Desalt）').font=inp_font
    for col in (19,20): wo.cell(row=i,column=col).fill=yellow
    wo.cell(row=i,column=12).number_format='0.00'; wo.cell(row=i,column=10).number_format='0.0'; wo.cell(row=i,column=15).number_format='0.0'
    for col in (16,17,18): wo.cell(row=i,column=col).alignment=Alignment(horizontal='center')
    wo.cell(row=i,column=4).font=Font(name='Consolas',size=10)
last=r0+len(W)-1
wo.cell(row=last+2,column=2,value='合计').font=bold
wo.cell(row=last+2,column=5,value='寡核苷酸数').font=bold
wo.cell(row=last+2,column=6,value=f'=COUNTA(B{r0}:B{last})').font=bold
for col,nm in ((16,'12 条池'),(17,'16 条池'),(18,'24 条池')):
    L=get_column_letter(col); wo.cell(row=last+2,column=col,value=f'=COUNTIF({L}{r0}:{L}{last},"✓")').font=bold; wo.cell(row=last+2,column=col).alignment=Alignment(horizontal='center')
    wo.cell(row=last+3,column=col,value=f'=SUMIF({L}{r0}:{L}{last},"✓",M{r0}:M{last})').font=bold; wo.cell(row=last+3,column=col).alignment=Alignment(horizontal='center')
wo.cell(row=last+2,column=15,value='各池寡核苷酸数').font=bold; wo.cell(row=last+3,column=15,value='各池展开序列总数').font=bold
wo.cell(row=last+5,column=2,value='黄色底 = 请确认或修改的订购设置（合成规模 25 nmol 和标准脱盐是占位默认值，不是推荐）。').font=f(italic=True,size=9)
for j,w in enumerate([6,24,26,28,8,9,9,9,12,9,14,8,9,13,14,8,8,8,10,20,8,40],1): wo.column_dimensions[get_column_letter(j)].width=w
wo.row_dimensions[1].height=45; wo.freeze_panes='E2'
# ---------- 池工作表
def pool_sheet(title,lv):
    wp=wb.create_sheet(title)
    wp['A1']=f'{title}配制表（100 µL 的 10× 混合液）'; wp['A1'].font=f(bold=True,size=12)
    wp['A3']='母液浓度 µM'; wp['B3']=100; wp['A4']='10× 混合液体积 µL'; wp['B4']=100; wp['A5']='每个位点在 10× 混合液里的浓度 µM'; wp['B5']=2.5
    wp['A6']='每个位点的母液体积 µL'; wp['B6']='=B4*B5/B3'
    wp['C5']='终浓度 = 该值 ÷ 10 = 0.25 µM（250 nM）'; wp['C6']='公式：体积 × 浓度 ÷ 母液浓度'
    for r in (3,4,5): wp.cell(row=r,column=2).font=inp_font; wp.cell(row=r,column=2).fill=yellow
    for r in (3,4,5,6): wp.cell(row=r,column=1).font=bold
    wp['B6'].font=base; wp['C5'].font=f(italic=True,size=9); wp['C6'].font=f(italic=True,size=9)
    h=['订购名称（ASCII）','序列 5′→3′','引物位点','展开序列数','位点内比例（输入）','加入体积 µL','位点内比例合计','核对']
    for j,t in enumerate(h,1):
        c=wp.cell(row=8,column=j,value=t); c.font=hdr_font; c.fill=hdr_fill; c.alignment=center; c.border=border
    col=f'uL_{lv}'; S=W[W[col].notna()].copy(); S['share']=[sh[lv].get(o) for o in S.order]
    start=9; end=start+len(S)-1
    for i,r in enumerate(S.itertuples(),start):
        wp.cell(row=i,column=1,value=r.name_ascii); wp.cell(row=i,column=2,value=r.seq); wp.cell(row=i,column=3,value=r.slot); wp.cell(row=i,column=4,value=int(r.expansions))
        c=wp.cell(row=i,column=5,value=float(round(r.share,3))); c.font=inp_font; c.number_format='0.0%'
        wp.cell(row=i,column=6,value=f'=E{i}*$B$6'); wp.cell(row=i,column=6).number_format='0.00'
        wp.cell(row=i,column=7,value=f'=SUMIF($C${start}:$C${end},C{i},$E${start}:$E${end})'); wp.cell(row=i,column=7).number_format='0.0%'
        wp.cell(row=i,column=8,value=f'=IF(ABS(G{i}-1)<0.002,"OK","请检查")')
        for j in range(1,9):
            c=wp.cell(row=i,column=j); c.border=border
            if j!=5: c.font=base
        wp.cell(row=i,column=2).font=Font(name='Consolas',size=10)
    wp.cell(row=end+2,column=1,value='合计').font=bold
    wp.cell(row=end+2,column=4,value=f'=SUM(D{start}:D{end})').font=bold; wp.cell(row=end+2,column=3,value=f'=COUNTA(A{start}:A{end})').font=bold
    wp.cell(row=end+2,column=6,value=f'=SUM(F{start}:F{end})').font=bold; wp.cell(row=end+2,column=6).number_format='0.00'
    wp.cell(row=end+3,column=1,value='加水 µL').font=bold; wp.cell(row=end+3,column=6,value=f'=B4-F{end+2}').font=bold; wp.cell(row=end+3,column=6).number_format='0.00'
    wp.cell(row=end+4,column=1,value='总体积 µL').font=bold; wp.cell(row=end+4,column=6,value=f'=F{end+2}+F{end+3}').font=bold; wp.cell(row=end+4,column=6).number_format='0.00'
    wp.cell(row=end+2,column=2,value='← 寡核苷酸数（C 列）、展开序列总数（D 列）、母液总体积（F 列）').font=f(italic=True,size=9)
    wp.cell(row=end+6,column=1,value='用法：每 25 µL PCR 加 2.5 µL 10× 混合液。比例列是输入（蓝色），修改后体积自动更新；同位点内比例合计应为 100%。').font=f(italic=True,size=9)
    for j,w in enumerate([26,28,10,10,14,12,14,10],1): wp.column_dimensions[get_column_letter(j)].width=w
    wp.column_dimensions['A'].width=34; wp.row_dimensions[8].height=32; wp.freeze_panes='A9'
    return wp,end
ENDS={}
for t,lv in (('12 条池','12'),('16 条池','16'),('24 条池','24')): ENDS[t]=pool_sheet(t,lv)[1]
# ---------- 池比较
wc=wb.create_sheet('池比较')
wc['A1']='三个池的评估摘要（序列层面的计算，非实测）'; wc['A1'].font=f(bold=True,size=12)
hh=['指标','12 条池','16 条池','24 条池','说明 / 来源']
for j,t in enumerate(hh,1):
    c=wc.cell(row=3,column=j,value=t); c.font=hdr_font; c.fill=hdr_fill; c.alignment=center; c.border=border
rows=[('寡核苷酸数',f"='12 条池'!C{ENDS['12 条池']+2}",f"='16 条池'!C{ENDS['16 条池']+2}",f"='24 条池'!C{ENDS['24 条池']+2}",'公式：来自各池工作表的合计行'),
('展开序列总数',f"='12 条池'!D{ENDS['12 条池']+2}",f"='16 条池'!D{ENDS['16 条池']+2}",f"='24 条池'!D{ENDS['24 条池']+2}",'公式：来自各池工作表的合计行'),
('综合得分 J',0.8815,0.9228,0.9319,'三项等权平均：属准确率（完整数据库）、扩增子覆盖（主要门平均）、支原体至少扩出 3 个扩增子；留出集 2,196 条序列'),
('J 的 95% 区间下限',0.8696,0.9150,0.9247,'对样本序列 400 次自助重抽样'),
('J 的 95% 区间上限',0.8913,0.9301,0.9385,'同上'),
('属准确率（完整数据库）',0.9383,0.9394,0.9413,'属级最近邻，样本只有扩出的扩增子'),
('扩增子覆盖（主要门平均）',0.8480,0.8969,0.9055,'Greengenes 与 SILVA 主要门平均的平均'),
('支原体至少扩出 3 个扩增子',0.8581,0.9323,0.9489,'缺数据的序列按未扩出算，数值偏保守'),
('6 个扩增子全部扩出，Greengenes',0.533,0.665,0.704,''),
('6 个扩增子全部扩出，SILVA',0.412,0.545,0.576,''),
('平均 Tm 最低 °C',54.8,56.1,54.8,'primer3，50 mM Na⁺、2 mM Mg²⁺、0.2 mM dNTP、250 nM 引物'),
('平均 Tm 最高 °C',66.9,66.9,66.9,''),
('严重二聚体数',0,0,0,'池内含展开序列两两检查（ΔG ≤ −9 kcal/mol 或 3′ 端 ≤ −6）'),
('人基因组 ≤2 错配位点',2849,702,3125,'hg19；16 条池低是因为用了 V1_f，不是条数'),
('人基因组 ≤3 错配潜在产物（80–1500 bp）',50,3,60,'同上'),]
for i,r in enumerate(rows,4):
    for j,v in enumerate(r,1):
        c=wc.cell(row=i,column=j,value=v); c.font=base; c.border=border; c.alignment=Alignment(vertical='top',wrap_text=(j==5))
        if j in (2,3,4) and isinstance(v,float): c.number_format='0.0%'
    if r[0] in ('综合得分 J','J 的 95% 区间下限','J 的 95% 区间上限'):
        for j in (2,3,4): wc.cell(row=i,column=j).number_format='0.0000'
    if r[0].startswith('平均 Tm'):
        for j in (2,3,4): wc.cell(row=i,column=j).number_format='0.0'
    if r[0] in ('人基因组 ≤2 错配位点',):
        for j in (2,3,4): wc.cell(row=i,column=j).number_format='#,##0'
for j,w in enumerate([36,12,12,12,70],1): wc.column_dimensions[get_column_letter(j)].width=w
wc.cell(row=len(rows)+6,column=1,value='以上评估数值是计算结果（硬编码，来源：primer_order_list_levels.html 和 docs/primer_design/lvl4_*.csv），不会随配制表变化；前两行是公式。').font=f(italic=True,size=9)
wb.calculation.fullCalcOnLoad=True
out='/home/user/16S_5R/docs/primer_design/primer_order_12_16_24.xlsx'; wb.save(out); print(out)
